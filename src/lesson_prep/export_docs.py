from __future__ import annotations

import io
import re
from typing import Any


def _as_list(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    return [str(x).strip() for x in value if str(x).strip()]


def _safe_filename(title: str) -> str:
    name = re.sub(r'[\\/:*?"<>|]+', "_", title).strip() or "备课结果"
    return name[:60]


def build_export_basename(lesson_input: dict[str, Any], module: str = "all") -> str:
    grade = str(lesson_input.get("grade") or "")
    title = str(lesson_input.get("lesson_title") or "备课结果")
    labels = {
        "all": "全部",
        "curriculum": "课标解读",
        "plan": "教案",
        "exercises": "习题卷",
        "slides": "课件大纲",
        "board": "板书",
    }
    suffix = labels.get(module, module)
    return _safe_filename(f"{grade}-{title}-{suffix}".strip("-"))


def _lesson_info_section(lesson_input: dict[str, Any]) -> tuple[str, list[str]]:
    return (
        "课时信息",
        [
            f"学段：{lesson_input.get('stage', '')}",
            f"年级：{lesson_input.get('grade', '')}",
            f"学科：{lesson_input.get('subject', '')}",
            f"教材：{lesson_input.get('textbook_version', '')}",
            f"单元：{lesson_input.get('unit', '')}",
            f"课题：{lesson_input.get('lesson_title', '')}",
            f"课时：{lesson_input.get('duration_minutes', '')} 分钟",
        ],
    )


def _collect_sections(
    lesson_input: dict[str, Any],
    result: dict[str, Any],
    module: str = "all",
) -> list[tuple[str, list[str]]]:
    curriculum = result.get("curriculum_analysis") or {}
    plan = result.get("lesson_plan") or {}
    exercises = result.get("exercise_paper") or {}
    slides = result.get("slides") or {}
    board = result.get("blackboard") or {}

    sections: list[tuple[str, list[str]]] = [_lesson_info_section(lesson_input)]

    if module in {"all", "curriculum"}:
        sections.extend(
            [
                ("核心素养", _as_list(curriculum.get("core_competencies"))),
                ("学业要求", _as_list(curriculum.get("academic_requirements"))),
                ("内容要点", _as_list(curriculum.get("content_points"))),
                ("教学提示", _as_list(curriculum.get("teaching_tips_from_standard"))),
            ]
        )

    if module in {"all", "plan"}:
        sections.extend(
            [
                ("教学目标", _as_list(plan.get("teaching_objectives"))),
                ("重点", _as_list(plan.get("key_points"))),
                ("难点", _as_list(plan.get("difficult_points"))),
                ("练习意图", _as_list(plan.get("practice_intents"))),
            ]
        )
        stages = plan.get("stages") if isinstance(plan.get("stages"), list) else []
        stage_lines: list[str] = []
        for stage in stages:
            if not isinstance(stage, dict):
                continue
            stage_lines.append(
                f"【{stage.get('name', '环节')}】{stage.get('duration_minutes', '')}分钟 | "
                f"师：{stage.get('teacher_activity', '')} | "
                f"生：{stage.get('student_activity', '')} | "
                f"目的：{stage.get('purpose', '')}"
            )
        sections.append(("环节设计", stage_lines))

    if module in {"all", "exercises"}:
        dist = exercises.get("difficulty_distribution") or {}
        sections.append(
            (
                "组卷说明",
                [
                    f"标题：{exercises.get('title', '')}",
                    f"总分：{exercises.get('total_score', '')}",
                    f"建议用时：{exercises.get('time_limit_minutes', '')} 分钟",
                    f"难度分布：易{dist.get('easy', 0)} / 中{dist.get('medium', 0)} / 难{dist.get('hard', 0)}",
                    f"知识点覆盖：{'、'.join(_as_list(exercises.get('knowledge_coverage')))}",
                    f"设计说明：{exercises.get('design_notes', '')}",
                ],
            )
        )
        items = exercises.get("items") if isinstance(exercises.get("items"), list) else []
        item_lines: list[str] = []
        for item in items:
            if not isinstance(item, dict):
                continue
            opts = "；".join(_as_list(item.get("options")))
            item_lines.append(
                f"第{item.get('index', '')}题[{item.get('difficulty', '')}/{item.get('question_type', '')}/"
                f"{item.get('score', '')}分] {item.get('stem', '')}"
                + (f" 选项：{opts}" if opts else "")
                + f" 答案：{item.get('answer', '')} 解析：{item.get('analysis', '')}"
            )
        sections.append(("题目", item_lines))

    if module in {"all", "slides"}:
        pages = slides.get("pages") if isinstance(slides.get("pages"), list) else []
        page_lines: list[str] = []
        if slides.get("design_notes"):
            page_lines.append(f"设计说明：{slides.get('design_notes')}")
        for page in pages:
            if not isinstance(page, dict):
                continue
            bullets = "；".join(_as_list(page.get("bullets")))
            page_lines.append(f"P{page.get('index', '')} {page.get('title', '')}：{bullets}")
        sections.append(("课件大纲", page_lines))

    if module in {"all", "board"}:
        board_lines = [f"布局：{board.get('layout', 'main_side')}"]
        main_board = board.get("main_board") if isinstance(board.get("main_board"), list) else []
        for item in sorted(
            main_board,
            key=lambda x: int((x or {}).get("order") or 0) if isinstance(x, dict) else 0,
        ):
            if isinstance(item, dict):
                board_lines.append(f"主板书：{item.get('text', '')}")
        side_board = board.get("side_board") if isinstance(board.get("side_board"), list) else []
        for item in side_board:
            if isinstance(item, dict):
                board_lines.append(f"副板书：{item.get('text', '')}")
        board_lines.extend([f"书写顺序：{x}" for x in _as_list(board.get("writing_sequence"))])
        board_lines.extend([f"关键句：{x}" for x in _as_list(board.get("key_sentences"))])
        sections.append(("板书设计", board_lines))

    return sections


MODULE_TITLES = {
    "all": "备课结果",
    "curriculum": "课标解读",
    "plan": "教案",
    "exercises": "习题卷",
    "slides": "课件大纲",
    "board": "板书设计",
}


def export_docx(lesson_input: dict[str, Any], result: dict[str, Any], module: str = "all") -> bytes:
    from docx import Document

    doc = Document()
    lesson = f"{lesson_input.get('grade', '')} {lesson_input.get('lesson_title', '')}".strip()
    module_title = MODULE_TITLES.get(module, "备课结果")
    doc.add_heading(f"{lesson} · {module_title}".strip(" ·") or module_title, level=0)
    doc.add_paragraph("智能备课教研工作台 · 分模块导出")

    for heading, lines in _collect_sections(lesson_input, result, module=module):
        if not lines:
            continue
        doc.add_heading(heading, level=1)
        for line in lines:
            doc.add_paragraph(line, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_pdf(lesson_input: dict[str, Any], result: dict[str, Any], module: str = "all") -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    font_name = _register_cn_font()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CNTitle",
        parent=styles["Title"],
        fontName=font_name,
        fontSize=18,
        leading=24,
        spaceAfter=12,
    )
    h_style = ParagraphStyle(
        "CNH",
        parent=styles["Heading2"],
        fontName=font_name,
        fontSize=13,
        leading=18,
        spaceBefore=10,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        "CNBody",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=10.5,
        leading=16,
        spaceAfter=3,
    )

    story = []
    lesson = f"{lesson_input.get('grade', '')} {lesson_input.get('lesson_title', '')}".strip()
    module_title = MODULE_TITLES.get(module, "备课结果")
    title = f"{lesson} · {module_title}".strip(" ·") or module_title
    story.append(Paragraph(_escape(title), title_style))
    story.append(Paragraph(_escape("智能备课教研工作台 · 分模块导出"), body_style))
    story.append(Spacer(1, 8))

    for heading, lines in _collect_sections(lesson_input, result, module=module):
        if not lines:
            continue
        story.append(Paragraph(_escape(heading), h_style))
        for line in lines:
            story.append(Paragraph(_escape(f"• {line}"), body_style))

    doc.build(story)
    return buf.getvalue()


def export_pptx(lesson_input: dict[str, Any], result: dict[str, Any], module: str = "slides") -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 与工作台一致的青绿配色
    colors = {
        "bg": RGBColor(0xF4, 0xF7, 0xF6),
        "panel": RGBColor(0xFF, 0xFF, 0xFF),
        "ink": RGBColor(0x17, 0x33, 0x3A),
        "slate": RGBColor(0x24, 0x54, 0x5A),
        "teal": RGBColor(0x2F, 0x6F, 0x6A),
        "teal_soft": RGBColor(0xD9, 0xE8, 0xE5),
        "amber": RGBColor(0xC4, 0x8A, 0x2A),
        "muted": RGBColor(0x5C, 0x72, 0x76),
        "line": RGBColor(0xC9, 0xD7, 0xD4),
        "white": RGBColor(0xFF, 0xFF, 0xFF),
    }

    title = str(lesson_input.get("lesson_title") or "课件大纲")
    grade = str(lesson_input.get("grade") or "")
    subject = str(lesson_input.get("subject") or "")
    unit = str(lesson_input.get("unit") or "")
    version = str(lesson_input.get("textbook_version") or "")
    duration = lesson_input.get("duration_minutes") or ""
    module_title = MODULE_TITLES.get(module, "课件大纲")

    _add_styled_cover(
        prs,
        colors,
        title=title,
        subtitle=f"{grade} · {subject} · {module_title}".strip(" ·"),
        meta="  |  ".join([x for x in [version, unit, f"{duration} 分钟" if duration else ""] if x]),
    )

    plan = result.get("lesson_plan") or {}
    slides = result.get("slides") or {}
    board = result.get("blackboard") or {}

    include_plan = module in {"all", "plan"}
    include_exercises = module in {"all", "exercises"}
    include_slides = module in {"all", "slides"}
    include_board = module in {"all", "board"}
    include_curriculum = module in {"all", "curriculum"}
    exercises = result.get("exercise_paper") or {}

    if include_curriculum:
        curriculum = result.get("curriculum_analysis") or {}
        _add_three_column_slide(
            prs,
            colors,
            page_title="课标解读要点",
            columns=[
                ("核心素养", _as_list(curriculum.get("core_competencies")) or ["（待补充）"]),
                ("学业要求", _as_list(curriculum.get("academic_requirements")) or ["（待补充）"]),
                ("内容要点", _as_list(curriculum.get("content_points")) or ["（待补充）"]),
            ],
            footer=title,
        )

    if include_plan:
        objectives = _as_list(plan.get("teaching_objectives"))
        key_points = _as_list(plan.get("key_points"))
        difficult = _as_list(plan.get("difficult_points"))
        if objectives or key_points or difficult:
            _add_three_column_slide(
                prs,
                colors,
                page_title="本节学习目标",
                columns=[
                    ("教学目标", objectives or ["（待补充）"]),
                    ("重点", key_points or ["（待补充）"]),
                    ("难点", difficult or ["（待补充）"]),
                ],
                footer=title,
            )
        stages = plan.get("stages") if isinstance(plan.get("stages"), list) else []
        for stage in stages:
            if not isinstance(stage, dict):
                continue
            _add_content_slide(
                prs,
                colors,
                eyebrow="教学环节",
                page_title=str(stage.get("name") or "环节"),
                bullets=[
                    f"时长：{stage.get('duration_minutes', '')} 分钟",
                    f"教师活动：{stage.get('teacher_activity', '')}",
                    f"学生活动：{stage.get('student_activity', '')}",
                    f"目的：{stage.get('purpose', '')}",
                ],
                callout="",
                tags=[],
                footer=title,
            )

    if include_exercises:
        items = exercises.get("items") if isinstance(exercises.get("items"), list) else []
        for item in items:
            if not isinstance(item, dict):
                continue
            bullets = [
                f"题型：{item.get('question_type', '')}　难度：{item.get('difficulty', '')}　"
                f"{item.get('score', '')}分",
                f"知识点：{item.get('knowledge_point', '')}",
                f"题目：{item.get('stem', '')}",
            ]
            opts = _as_list(item.get("options"))
            if opts:
                bullets.extend(opts)
            bullets.append(f"答案：{item.get('answer', '')}")
            if item.get("analysis"):
                bullets.append(f"解析：{item.get('analysis', '')}")
            _add_content_slide(
                prs,
                colors,
                eyebrow=f"习题 第{item.get('index', '')}题",
                page_title=str(exercises.get("title") or "随堂练习"),
                bullets=bullets,
                callout="",
                tags=_as_list(exercises.get("knowledge_coverage"))[:3],
                footer=title,
            )

    if include_slides:
        pages = slides.get("pages") if isinstance(slides.get("pages"), list) else []
        if pages:
            for page in pages:
                if not isinstance(page, dict):
                    continue
                bullets = _as_list(page.get("bullets"))
                interaction = str(page.get("interaction") or "").strip()
                keywords = _as_list(page.get("visual_keywords"))
                _add_content_slide(
                    prs,
                    colors,
                    eyebrow=f"课件 P{page.get('index', '')}",
                    page_title=str(page.get("title") or "内容页"),
                    bullets=bullets or ["（本页暂无要点）"],
                    callout=f"互动：{interaction}" if interaction else "",
                    tags=keywords,
                    footer=title,
                )
        elif module == "slides":
            # 仅导出课件但暂无 pages 时，用教案环节兜底
            stages = plan.get("stages") if isinstance(plan.get("stages"), list) else []
            for stage in stages:
                if not isinstance(stage, dict):
                    continue
                _add_content_slide(
                    prs,
                    colors,
                    eyebrow="教学环节",
                    page_title=str(stage.get("name") or "环节"),
                    bullets=[
                        f"时长：{stage.get('duration_minutes', '')} 分钟",
                        f"教师活动：{stage.get('teacher_activity', '')}",
                        f"学生活动：{stage.get('student_activity', '')}",
                        f"目的：{stage.get('purpose', '')}",
                    ],
                    callout="",
                    tags=[],
                    footer=title,
                )

    if include_board:
        main_board = [
            str(x.get("text") or "")
            for x in (board.get("main_board") or [])
            if isinstance(x, dict) and str(x.get("text") or "").strip()
        ]
        side_board = [
            str(x.get("text") or "")
            for x in (board.get("side_board") or [])
            if isinstance(x, dict) and str(x.get("text") or "").strip()
        ]
        if main_board or side_board:
            _add_two_column_slide(
                prs,
                colors,
                page_title="板书提示",
                left_title=f"主板书（{board.get('layout', 'main_side')}）",
                left_items=main_board or ["（无）"],
                right_title="副板书 / 易错提醒",
                right_items=side_board or ["（无）"],
                footer=title,
            )

    _add_styled_ending(prs, colors, title=title)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _fill_solid(shape, color) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _add_bg(slide, colors) -> None:
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    _fill_solid(bg, colors["bg"])
    # 左侧色带
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), Inches(7.5))
    _fill_solid(bar, colors["teal"])
    # 顶部淡色条
    top = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.18), 0, Inches(13.153), Inches(0.08))
    _fill_solid(top, colors["teal_soft"])


def _add_footer(slide, colors, text: str) -> None:
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    box = slide.shapes.add_textbox(Inches(0.55), Inches(7.05), Inches(12.2), Inches(0.3))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = f"Lesson Atelier  ·  {text}"
    run.font.size = Pt(11)
    run.font.name = "微软雅黑"
    run.font.color.rgb = colors["muted"]


def _add_styled_cover(prs, colors, *, title: str, subtitle: str, meta: str) -> None:
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    _add_bg(slide, colors)

    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.5), Inches(11.5), Inches(4.3))
    _fill_solid(panel, colors["panel"])
    panel.line.color.rgb = colors["line"]

    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.5), Inches(0.18), Inches(4.3))
    _fill_solid(accent, colors["teal"])

    brand = slide.shapes.add_textbox(Inches(1.5), Inches(1.85), Inches(10), Inches(0.4))
    p = brand.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "LESSON ATELIER  ·  智能备课"
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = colors["teal"]
    run.font.name = "微软雅黑"

    title_box = slide.shapes.add_textbox(Inches(1.5), Inches(2.5), Inches(10.2), Inches(1.6))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.size = Pt(40)
    run.font.bold = True
    run.font.color.rgb = colors["ink"]
    run.font.name = "微软雅黑"

    sub = slide.shapes.add_textbox(Inches(1.5), Inches(4.2), Inches(10.2), Inches(0.5))
    p = sub.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = subtitle or "课堂课件大纲"
    run.font.size = Pt(20)
    run.font.color.rgb = colors["slate"]
    run.font.name = "微软雅黑"

    meta_box = slide.shapes.add_textbox(Inches(1.5), Inches(4.9), Inches(10.2), Inches(0.4))
    p = meta_box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = meta
    run.font.size = Pt(14)
    run.font.color.rgb = colors["muted"]
    run.font.name = "微软雅黑"

    for i, c in enumerate([colors["teal"], colors["amber"], colors["teal_soft"]]):
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(11.3 + i * 0.28),
            Inches(5.35),
            Inches(0.18),
            Inches(0.18),
        )
        _fill_solid(dot, c)


def _add_styled_ending(prs, colors, *, title: str) -> None:
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, colors)
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.2), Inches(2.2), Inches(8.9), Inches(3.0))
    _fill_solid(panel, colors["panel"])
    panel.line.color.rgb = colors["line"]

    box = slide.shapes.add_textbox(Inches(2.6), Inches(2.7), Inches(8.1), Inches(2.0))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "谢谢聆听"
    run.font.size = Pt(36)
    run.font.bold = True
    run.font.color.rgb = colors["ink"]
    run.font.name = "微软雅黑"

    p2 = tf.add_paragraph()
    run2 = p2.add_run()
    run2.text = f"\n{title}\n请带着问题走进下一环节"
    run2.font.size = Pt(16)
    run2.font.color.rgb = colors["slate"]
    run2.font.name = "微软雅黑"


def _add_content_slide(
    prs,
    colors,
    *,
    eyebrow: str,
    page_title: str,
    bullets: list[str],
    callout: str,
    tags: list[str],
    footer: str,
) -> None:
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, colors)

    eyebrow_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.35), Inches(11.5), Inches(0.35))
    p = eyebrow_box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = eyebrow
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = colors["teal"]
    run.font.name = "微软雅黑"

    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.7), Inches(11.8), Inches(0.7))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = page_title
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = colors["ink"]
    run.font.name = "微软雅黑"

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(1.55), Inches(12.0), Inches(4.7))
    _fill_solid(card, colors["panel"])
    card.line.color.rgb = colors["line"]

    body = slide.shapes.add_textbox(Inches(1.1), Inches(1.85), Inches(11.2), Inches(3.5))
    tf = body.text_frame
    tf.word_wrap = True
    for i, line in enumerate(bullets[:8]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(10)
        run = p.add_run()
        run.text = f"▸  {line}"
        run.font.size = Pt(18)
        run.font.color.rgb = colors["ink"]
        run.font.name = "微软雅黑"

    y = 5.55
    if callout:
        tip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1), Inches(y), Inches(7.2), Inches(0.45))
        _fill_solid(tip, colors["teal_soft"])
        tip.line.fill.background()
        tip_box = slide.shapes.add_textbox(Inches(1.25), Inches(y + 0.05), Inches(6.9), Inches(0.35))
        p = tip_box.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = callout
        run.font.size = Pt(12)
        run.font.color.rgb = colors["slate"]
        run.font.name = "微软雅黑"

    if tags:
        tag_box = slide.shapes.add_textbox(Inches(8.5), Inches(y + 0.05), Inches(3.8), Inches(0.35))
        p = tag_box.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = " · ".join(tags[:4])
        run.font.size = Pt(11)
        run.font.color.rgb = colors["amber"]
        run.font.name = "微软雅黑"

    _add_footer(slide, colors, footer)


def _add_three_column_slide(prs, colors, *, page_title: str, columns: list[tuple[str, list[str]]], footer: str) -> None:
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, colors)

    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(12), Inches(0.6))
    p = title_box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = page_title
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = colors["ink"]
    run.font.name = "微软雅黑"

    width = 3.8
    gap = 0.25
    start_x = 0.7
    for i, (col_title, items) in enumerate(columns[:3]):
        x = start_x + i * (width + gap)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.3), Inches(width), Inches(5.2))
        _fill_solid(card, colors["panel"])
        card.line.color.rgb = colors["line"]
        head = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(1.3), Inches(width), Inches(0.55))
        _fill_solid(head, colors["teal"] if i != 2 else colors["amber"])

        hbox = slide.shapes.add_textbox(Inches(x + 0.2), Inches(1.38), Inches(width - 0.35), Inches(0.4))
        p = hbox.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = col_title
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = colors["white"]
        run.font.name = "微软雅黑"

        body = slide.shapes.add_textbox(Inches(x + 0.25), Inches(2.1), Inches(width - 0.45), Inches(4.1))
        tf = body.text_frame
        tf.word_wrap = True
        for j, item in enumerate(items[:7]):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.space_after = Pt(8)
            run = p.add_run()
            run.text = f"• {item}"
            run.font.size = Pt(14)
            run.font.color.rgb = colors["ink"]
            run.font.name = "微软雅黑"

    _add_footer(slide, colors, footer)


def _add_two_column_slide(
    prs,
    colors,
    *,
    page_title: str,
    left_title: str,
    left_items: list[str],
    right_title: str,
    right_items: list[str],
    footer: str,
) -> None:
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, colors)

    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(12), Inches(0.6))
    p = title_box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = page_title
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = colors["ink"]
    run.font.name = "微软雅黑"

    for x, w, head_color, col_title, items in [
        (0.7, 7.4, colors["teal"], left_title, left_items),
        (8.35, 4.3, colors["amber"], right_title, right_items),
    ]:
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.3), Inches(w), Inches(5.2))
        _fill_solid(card, colors["panel"])
        card.line.color.rgb = colors["line"]
        head = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(1.3), Inches(w), Inches(0.55))
        _fill_solid(head, head_color)
        hbox = slide.shapes.add_textbox(Inches(x + 0.25), Inches(1.38), Inches(w - 0.4), Inches(0.4))
        p = hbox.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = col_title
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = colors["white"]
        run.font.name = "微软雅黑"

        body = slide.shapes.add_textbox(Inches(x + 0.3), Inches(2.15), Inches(w - 0.55), Inches(4.0))
        tf = body.text_frame
        tf.word_wrap = True
        for j, item in enumerate(items[:10]):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.space_after = Pt(8)
            run = p.add_run()
            run.text = f"• {item}"
            run.font.size = Pt(15)
            run.font.color.rgb = colors["ink"]
            run.font.name = "微软雅黑"

    _add_footer(slide, colors, footer)


def _escape(text: str) -> str:
    return (
        str(text)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def _register_cn_font() -> str:
    """注册系统中文字体，保证 PDF 中文可显示。"""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [
        ("SimSun", r"C:\Windows\Fonts\simsun.ttc"),
        ("SimHei", r"C:\Windows\Fonts\simhei.ttf"),
        ("MicrosoftYaHei", r"C:\Windows\Fonts\msyh.ttc"),
        ("NotoSansSC", r"C:\Windows\Fonts\NotoSansSC-Regular.otf"),
    ]
    for name, path in candidates:
        try:
            pdfmetrics.registerFont(TTFont(name, path))
            return name
        except Exception:  # noqa: BLE001
            continue
    return "Helvetica"
